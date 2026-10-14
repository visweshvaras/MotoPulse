#!/usr/bin/env python3
"""
Generate a professionally styled Excel spreadsheet (.xlsx) and CSV (.csv)
for the ESP32-S3 LM358 Microphone Preamplifier breadboard shopping list.
"""

import os
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = "/home/vishwa/Documents/design"
XLSX_PATH = os.path.join(BASE_DIR, "breadboard_shopping_list.xlsx")
CSV_PATH = os.path.join(BASE_DIR, "breadboard_shopping_list.csv")

data = [
    # Headers: Category, Ref Designator, Component Name, Value / Specification, Package / Form Factor, Marking / Color Code, Qty, Purpose / Buying Notes
    ["Active IC", "U1", "LM358 Dual Operational Amplifier", "Single supply 3.0V - 32V, Low power", "DIP-8 (Through-Hole)", "LM358N or LM358P", 1, "Core preamplifier. Must be DIP-8 package with pins for breadboard."],
    ["Microphone", "MK1 / J2", "Electret Condenser Microphone Capsule", "Standard omnidirectional, 2V - 10V operating", "2-Pin Radial Leaded (6mm or 9mm)", "Terminal connected to metal case is GND (-)", 1, "Acoustic audio pickup. One pin connects to case (GND)."],
    ["Resistor", "R1", "Metal Film / Carbon Film Resistor", "10 kΩ, 1/4W (0.25W), 5% or 1%", "Axial Through-Hole", "Brown - Black - Orange - Gold", 1, "Microphone bias pull-up to +3.3V (supplies ~0.3mA to internal JFET)."],
    ["Resistor", "R2", "Metal Film / Carbon Film Resistor", "100 kΩ, 1/4W (0.25W), 5% or 1%", "Axial Through-Hole", "Brown - Black - Yellow - Gold", 1, "DC mid-rail bias voltage divider upper leg (pull-up to +3.3V)."],
    ["Resistor", "R3", "Metal Film / Carbon Film Resistor", "100 kΩ, 1/4W (0.25W), 5% or 1%", "Axial Through-Hole", "Brown - Black - Yellow - Gold", 1, "DC mid-rail bias voltage divider lower leg (pull-down to GND)."],
    ["Resistor", "R4", "Metal Film / Carbon Film Resistor", "100 kΩ, 1/4W (0.25W), 5% or 1%", "Axial Through-Hole", "Brown - Black - Yellow - Gold", 1, "Negative feedback resistor. Sets AC voltage gain (Av = 1 + R4/R5 = 101)."],
    ["Resistor", "R5", "Metal Film / Carbon Film Resistor", "1 kΩ, 1/4W (0.25W), 5% or 1%", "Axial Through-Hole", "Brown - Black - Red - Gold", 1, "Inverting input gain resistor to AC ground."],
    ["Resistor", "R6", "Metal Film / Carbon Film Resistor", "1 kΩ, 1/4W (0.25W), 5% or 1%", "Axial Through-Hole", "Brown - Black - Red - Gold", 1, "Output RC low-pass filter series resistor & ESP32 ADC current limiter."],
    ["Capacitor", "C1", "Multilayer Ceramic (MLCC) / Film Capacitor", "1 µF, 16V - 50V, Non-polarized", "Radial Leaded (2.54mm or 5mm pitch)", "Marked '105' or '1µF'", 1, "AC audio coupling. Blocks DC mic bias while passing voice audio."],
    ["Capacitor", "C2", "Aluminum Electrolytic Capacitor", "10 µF, 16V - 25V, Polarized", "Radial Leaded Through-Hole", "Marked '10µF', Stripe on (-) leg", 1, "Power supply bulk decoupling on breadboard +3.3V rail."],
    ["Capacitor", "C3", "Aluminum Electrolytic Capacitor", "10 µF, 16V - 25V, Polarized", "Radial Leaded Through-Hole", "Marked '10µF', Stripe on (-) leg", 1, "AC ground capacitor in series with R5. Keeps DC gain at unity (1.0)."],
    ["Capacitor", "C4", "Ceramic Disc / Multilayer Ceramic", "10 nF (0.01 µF), 50V, Non-polarized", "Radial Leaded (2.54mm pitch)", "Marked '103'", 1, "Anti-aliasing low-pass filter shunt capacitor (fc = 15.9 kHz)."],
    ["Capacitor", "C5", "Ceramic Disc / Multilayer Ceramic", "100 nF (0.1 µF), 50V, Non-polarized", "Radial Leaded (2.54mm pitch)", "Marked '104'", 1, "High-frequency power bypass capacitor placed directly at LM358 Pin 8."],
    ["Microcontroller", "ESP32-S3", "ESP32-S3 Development Board", "Xtensa Dual-Core LX7 @ 240MHz, 12-Bit ADC", "Standard DevKit with pre-soldered pin headers", "ESP32-S3-DevKitC-1 or compatible", 1, "Target processor. Reads analog microphone data on GPIO 1 (ADC1_CH0)."],
    ["Hardware", "CABLE", "USB Type-C Data Cable", "Standard USB-A to USB-C (or USB-C to USB-C)", "Flexible Cable", "Must support Data (not charge-only)", 1, "Powers the ESP32-S3 and streams ADC audio readings to PC."],
    ["Hardware", "BOARD", "Solderless Breadboard", "830 Tie-Points with Dual Power Rails", "Model MB-102 or standard full-size breadboard", "Clear coordinate markings (A-J, 1-63)", 1, "Base prototyping platform for assembling the components without soldering."],
    ["Hardware", "WIRES", "Male-to-Male Dupont Jumper Wires", "Standard 2.54mm pitch jumper wires", "Pack of 20 - 40 jumper wires", "Assorted colors (Red, Black, Green, etc.)", 1, "For connecting power rails, IC pins, and ESP32 GPIO connections."],
    ["Tools (Optional)", "METER", "Digital Multimeter", "Basic voltage (DC/AC) & continuity multimeter", "Handheld test instrument", "Includes test probes", 1, "Highly recommended to verify 1.65V DC operating point before coding."]
]

headers = [
    "Category",
    "Ref Designator",
    "Component Name",
    "Value / Specification",
    "Package / Form Factor",
    "Marking / Color Code",
    "Qty",
    "Circuit Purpose / Buying Notes"
]

def generate_csv():
    with open(CSV_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(data)
    print(f"CSV saved to {CSV_PATH}")

def generate_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Breadboard BOM & Shopping List"
    
    # Title Block
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = "ESP32-S3 Electret Microphone Preamplifier (LM358) — Breadboard Shopping List"
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    ws.merge_cells("A2:H2")
    sub_cell = ws["A2"]
    sub_cell.value = "Complete Bill of Materials for Solderless Breadboard Assembly | Gain: 40.1 dB (x101) | Target ADC: ESP32-S3 GPIO 1"
    sub_cell.font = Font(name="Calibri", size=11, italic=True, color="CBD5E1")
    sub_cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 24
    
    # Empty separator row
    ws.row_dimensions[3].height = 10
    
    # Header Row
    header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid") # Ocean Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    
    for col_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if h_text in ["Ref Designator", "Qty"] else "left", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[4].height = 28

    # Category Colors
    cat_colors = {
        "Active IC": "FEF3C7",      # Light Amber
        "Microphone": "E0F2FE",     # Light Sky
        "Resistor": "F1F5F9",       # Light Slate
        "Capacitor": "FEE2E2",      # Light Rose
        "Microcontroller": "DCFCE7",# Light Emerald
        "Hardware": "F3E8FF",       # Light Purple
        "Tools (Optional)": "F8FAFC"
    }

    # Data Rows
    for row_idx, row_data in enumerate(data, 5):
        category = row_data[0]
        row_fill_color = cat_colors.get(category, "FFFFFF")
        row_fill = PatternFill(start_color=row_fill_color, end_color=row_fill_color, fill_type="solid")
        
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10, bold=(col_idx == 2 or col_idx == 7))
            cell.border = thin_border
            cell.fill = row_fill
            
            # Alignments
            if col_idx in [2, 7]: # Ref, Qty
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=(col_idx == 8))
                
        ws.row_dimensions[row_idx].height = 24

    # Adjust Column Widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            # Skip merged title rows
            if cell.row in [1, 2, 3]:
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        # Apply padding
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    # Extra tuning for notes column
    ws.column_dimensions["H"].width = 50

    wb.save(XLSX_PATH)
    print(f"XLSX saved to {XLSX_PATH}")

if __name__ == "__main__":
    generate_csv()
    generate_xlsx()
